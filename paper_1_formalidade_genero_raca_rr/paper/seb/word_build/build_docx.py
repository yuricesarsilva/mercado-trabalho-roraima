"""Gera artigo_seb.docx a partir de artigo_seb.tex (II Semana de Economia Brasileira).

Uso: python build_docx.py   (a partir de paper/seb/word_build; lê ../artigo_seb.tex)
"""
import io, os, re, sys, subprocess
import pypandoc

PAPER = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(os.path.abspath(__file__)), '..')
WB = os.path.dirname(os.path.abspath(__file__))
PANDOC = pypandoc.get_pandoc_path()

tex = io.open(os.path.join(PAPER, 'artigo_seb.tex'), encoding='utf-8').read()
aux = io.open(os.path.join(PAPER, 'artigo_seb.aux'), encoding='utf-8').read()
labels = dict(re.findall(r'\\newlabel\{([^}]*)\}\{\{([^}]*)\}', aux))

# --- preâmbulo: pandoc só precisa de título/autores ------------------------
body = tex[tex.index(r'\begin{document}') + len(r'\begin{document}'):tex.index(r'\end{document}')]

# folha de rosto: reescrita com marcadores que o filtro Lua converte em estilos
cover_end = body.index(r'\section{Introdução}')
cover = body[:cover_end]
resumo = re.search(r'\\textbf\{Resumo\}(.*?)\\medskip', cover, re.S).group(1)
pc = re.search(r'\\textbf\{Palavras-chave:\}(.*?)\n\n', cover, re.S).group(1)
abstract = re.search(r'\\textbf\{Abstract\}(.*?)\\medskip', cover, re.S).group(1)
kw = re.search(r'\\textbf\{Keywords:\}(.*?)\n\n', cover, re.S).group(1)
jel = re.search(r'\\textbf\{Classificação JEL:\}(.*?)\n', cover, re.S).group(1)
authors = re.findall(r'\n([^\n\\]+?)\\footnote\{([^}]*)\}', cover)
title = re.search(r'\\bfseries (.*?)\\par', cover).group(1)

def clean(t):
    return re.sub(r'\\noindent\s*', '', t).strip()

new_cover = ['COVERTITLE ' + title, '']
for name, aff in authors:
    new_cover += ['COVERAUTHOR ' + name.strip() + r'\footnote{' + aff + '}', '']
new_cover += [
    'COVERHEAD Resumo', '', 'COVERTEXT ' + clean(resumo), '',
    r'COVERTEXT \textbf{Palavras-chave:} ' + clean(pc), '',
    'COVERHEAD Abstract', '', 'COVERTEXT ' + clean(abstract), '',
    r'COVERTEXT \textbf{Keywords:} ' + clean(kw), '',
    r'COVERTEXT \textbf{Classificação JEL:} ' + clean(jel), '',
    'PAGEBREAKHERE', '', '']
body = '\n'.join(new_cover) + body[cover_end:]

# --- referências cruzadas resolvidas pelo .aux -----------------------------
eqs = {}
def eq_tag(m):
    lab = m.group(1)
    return r'\qquad (' + labels[lab] + ')'
body = re.sub(r'\\label\{(eq:[^}]*)\}', eq_tag, body)
body = re.sub(r'\\eqref\{([^}]*)\}', lambda m: '(' + labels[m.group(1)] + ')', body)
body = re.sub(r'\\ref\{([^}]*)\}', lambda m: labels[m.group(1)], body)
body = body.replace('~', ' ')

# legendas numeradas como no PDF
def number_captions(env, prefix):
    global body
    out, pos = [], 0
    for m in re.finditer(r'\\begin\{' + env + r'\}.*?\\end\{' + env + r'\}', body, re.S):
        blk = m.group(0)
        lab = re.search(r'\\label\{([^}]*)\}', blk)
        num = labels[lab.group(1)] if lab else '?'
        blk = re.sub(r'\\caption\{', r'\\caption{' + prefix + ' ' + num + ' -- ', blk, count=1)
        # notas (minipage) viram parágrafo após o float
        note = re.search(r'\\begin\{minipage\}\{[^}]*\}(.*?)\\end\{minipage\}', blk, re.S)
        if note:
            blk = blk.replace(note.group(0), '')
            n = re.sub(r'\\vspace\{[^}]*\}|\\scriptsize|\\footnotesize|\\small', '', note.group(1)).strip()
            blk = blk + '\n\nFLOATNOTE ' + n + '\n'
        out.append(body[pos:m.start()]); out.append(blk); pos = m.end()
    out.append(body[pos:])
    body = ''.join(out)

number_captions('figure', 'Figura')
number_captions('table', 'Tabela')

# figuras em PNG
body = body.replace('../figures/', 'figures/')
body = re.sub(r'(figures/[^}]*?)\.pdf', r'\1.png', body)

# bibliografia: marcador na posição correta (antes do apêndice)
body = re.sub(r'\\begingroup\s*(?:\\small)?\\singlespacing\s*\\setlength\{\\bibsep\}\{[^}]*\}\s*'
              r'\\bibliographystyle\{[^}]*\}\s*\\bibliography\{[^}]*\}\s*\\endgroup',
              r'\\section{Referências}\n\nREFSHERE\n', body)
assert 'REFSHERE' in body
# o apêndice começa em página nova
body = body.replace(r'\section*{Apêndice A', 'PAGEBREAKHERE\n\n' + r'\section*{Apêndice A')

# comandos de layout sem efeito no Word
for cmd in [r'\appendix', r'\normalsize', r'\begingroup', r'\endgroup', r'\noindent',
            r'\singlespacing', r'\centering']:
    body = body.replace(cmd, '')
body = re.sub(r'\\renewcommand\{\\the[^}]*\}\{(?:[^{}]|\{[^{}]*\})*\}', '', body)
body = re.sub(r'\\setcounter\{[^}]*\}\{[^}]*\}', '', body)

doc = (r'\documentclass{article}' + '\n' + r'\usepackage{amsmath}' + '\n' +
       r'\begin{document}' + '\n' + body + '\n' + r'\end{document}' + '\n')
src = os.path.join(WB, 'seb_word.tex')
io.open(src, 'w', encoding='utf-8').write(doc)

out = os.path.join(PAPER, 'artigo_seb.docx')
cmd = [PANDOC, src, '-f', 'latex', '-t', 'docx', '-o', out,
       '--lua-filter', os.path.join(WB, 'seb.lua'),
       '--citeproc', '--bibliography', os.path.join(PAPER, '..', 'referencias.bib'),
       '--csl', os.path.join(WB, 'chicago.csl'),
       '-M', 'lang=pt-BR', '-M', 'link-citations=false',
       '--number-sections',
       '--resource-path', WB,
       '--reference-doc', os.path.join(WB, 'reference.docx')]
r = subprocess.run(cmd, capture_output=True, text=True, encoding='utf-8')
print(r.stdout, r.stderr)
assert r.returncode == 0

# --- pós-processamento: tabelas na largura do texto, fonte 9, espaço simples ---
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

TEXT_W = 16.0  # cm (21 - 2 x 2,5)
dx = Document(out)
for t in dx.tables:
    tblPr = t._tbl.tblPr
    w = tblPr.find(qn('w:tblW'))
    if w is None:
        w = OxmlElement('w:tblW'); tblPr.append(w)
    w.set(qn('w:type'), 'pct'); w.set(qn('w:w'), '5000')
    lay = tblPr.find(qn('w:tblLayout'))
    if lay is None:
        lay = OxmlElement('w:tblLayout'); tblPr.append(lay)
    lay.set(qn('w:type'), 'fixed')
    ncols = len(t._tbl.tblGrid.findall(qn('w:gridCol')))
    first = 0.28 * TEXT_W
    rest = (TEXT_W - first) / (ncols - 1)
    widths = [first] + [rest] * (ncols - 1)
    for gc, wd in zip(t._tbl.tblGrid.findall(qn('w:gridCol')), widths):
        gc.set(qn('w:w'), str(int(Cm(wd).twips)))
    for row in t.rows:
        for j, cell in enumerate(row.cells):
            tcW = cell._tc.get_or_add_tcPr().find(qn('w:tcW'))
            if tcW is not None:
                tcW.set(qn('w:type'), 'auto'); tcW.set(qn('w:w'), '0')
            for p in cell.paragraphs:
                pf = p.paragraph_format
                pf.first_line_indent = Cm(0); pf.left_indent = Cm(0)
                pf.space_before = pf.space_after = Pt(0); pf.line_spacing = 1.0
                pf.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size = Pt(9)
dx.save(out)
print('docx:', out)
